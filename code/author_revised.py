from pathlib import Path
import os, sys, json, re, shutil, zipfile
ROOT=Path(__file__).resolve().parents[2] if Path(__file__).parent.name=='code' else Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'research_work'/'vendor'))
os.environ['MPLCONFIGDIR']=str(Path(__file__).resolve().parent.parent/'.cache'/'matplotlib')
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from latex2mathml.converter import convert
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle,Circle,FancyArrowPatch

OUT=Path(__file__).resolve().parent.parent;DATA=OUT/'data';FIG=OUT/'figures'
import time
while True:
    try:
        if (FIG/'fig11_timing_ablation.png').exists() and len(json.loads((DATA/'timing_electrical_report.json').read_text())['electrical'])==4:break
    except (FileNotFoundError,json.JSONDecodeError):pass
    time.sleep(5)
extra=json.loads((DATA/'additional_report.json').read_text())
lhs=json.loads((DATA/'lhs_report.json').read_text())
fe=json.loads((DATA/'finite_element_report.json').read_text())
refine=json.loads((DATA/'refinement_hardware_report.json').read_text())
bench=json.loads((DATA/'timing_electrical_report.json').read_text())
observer_report=json.loads((DATA/'observer_refinement_report.json').read_text())
rbf_report=json.loads((DATA/'rbf_report.json').read_text())
report=json.loads((DATA/'report.json').read_text(encoding='utf-8'))
eval_n=report['evaluation_modes'];eval_dt=report['simulation_step']
mv=json.loads((DATA/'matlab_validation.json').read_text(encoding='utf-8')) if (DATA/'matlab_validation.json').exists() else {}
nom={(r['controller'],r['direction']):r for r in report['nominal']}
ex=nom['Neural filtered',1];en=nom['Neural filtered',-1];lx=nom['Thermal LQR',1];ln=nom['Thermal LQR',-1]
ox=nom['Oracle',1];on=nom['Oracle',-1];v=report['verification'];learn=report['learning']
mc=report['monte_carlo'];pairs=[]
for j in sorted({r['sample'] for r in mc}):
    a=next(r for r in mc if r['sample']==j and r['controller']=='Thermal LQR');b=next(r for r in mc if r['sample']==j and r['controller']=='Neural filtered');pairs.append((a,b))
reduction=lambda a,b:100*(1-b/a)
costratio=np.array([b['cost']/a['cost'] for a,b in pairs]);vratio=np.array([b['rms_vibration_um']/a['rms_vibration_um'] for a,b in pairs]);aratio=np.array([b['rms_attitude_arcsec']/a['rms_attitude_arcsec'] for a,b in pairs])

def schematic():
    fig,ax=plt.subplots(figsize=(7.4,4.4));ax.set_xlim(-.5,10.2);ax.set_ylim(0,6);ax.axis('off')
    ax.add_patch(Circle((1.05,4.4),.6,facecolor='#DCE5EF',edgecolor='black'));ax.text(1.05,4.4,'Rigid hub\nθ and τ',ha='center',va='center',fontsize=10)
    ax.add_patch(Rectangle((1.65,4.15),7.7,.35,facecolor='#EEF1F3',edgecolor='black'))
    for x in [1.8,3.3,5.6]:
        ax.add_patch(Rectangle((x,4.5),.65,.10,facecolor='#1768AF'));ax.add_patch(Rectangle((x,4.05),.65,.1,facecolor='#1768AF'))
    ax.text(5.9,3.77,'Piezoelectric patch pairs and elastic coordinates q',ha='center',fontsize=9)
    for x in [3.,4.,5.,6.,7.,8.]:ax.annotate('',xy=(x,4.8),xytext=(x-.3,5.6),arrowprops=dict(arrowstyle='->',color='#D99020',lw=1.6))
    ax.text(6,5.8,'Solar heat flux and photovoltaic conversion',ha='center',fontsize=10)
    ax.text(9.5,4.35,'w(L,t)',fontsize=9,rotation=90,va='center')
    boxes=[(0,1.5,2.1,'Two surface\nthermal states'),(2.6,1.5,2.0,'Moving thermal\nequilibrium'),(5.1,1.5,2.1,'LQR and learned\noptimal preview'),(7.7,1.5,2.3,'Lyapunov filter\nand input bounds')]
    for x,y,w,label in boxes:
        ax.add_patch(Rectangle((x,y),w,1.1,facecolor='white',edgecolor='#3B4652',lw=1.));ax.text(x+w/2,y+.55,label,ha='center',va='center',fontsize=9)
    for a,b in [(2.1,2.6),(4.6,5.1),(7.2,7.7)]:ax.annotate('',xy=(b,2.05),xytext=(a,2.05),arrowprops=dict(arrowstyle='->',lw=1.))
    ax.annotate('',xy=(8.8,3.98),xytext=(8.8,2.65),arrowprops=dict(arrowstyle='->',lw=1.));ax.text(9.05,3.1,'Voltage\nand torque',fontsize=8)
    ax.annotate('',xy=(6.1,1.4),xytext=(6.1,.65),arrowprops=dict(arrowstyle='->',lw=1.));ax.plot([1.05,-.3,-.3,6.1],[3.7,3.7,.65,.65],color='#46566A',lw=1.)
    ax.text(3.6,.35,'Retained states, surface temperatures and residual strain-rate damping',ha='center',fontsize=8.5)
    fig.tight_layout()
    for ext in ['png','pdf','svg']:fig.savefig(FIG/('fig01_architecture.'+ext),dpi=360,bbox_inches='tight')
    plt.close(fig)

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.10;normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    if name.startswith('Heading'):style.font.bold=True;style.paragraph_format.keep_with_next=True;style.paragraph_format.space_before=Pt(11);style.paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].font.bold=True
doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 2'].font.size=Pt(11.5)
doc.styles['Caption'].font.size=Pt(9.5);doc.styles['Caption'].paragraph_format.space_after=Pt(9)
for st in doc.styles:
    if st.type==1:
        st.font.color.rgb=RGBColor(0,0,0)
        if st.name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
            fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
            for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Times New Roman')
            for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+key),None)
        for node in st.element.xpath('.//w:pBdr'):node.getparent().remove(node)
doc.styles['Caption'].font.bold=False
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();r.font.size=Pt(9);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.append(field)
doc.core_properties.title='Thermal Preview Learning and Lyapunov Safeguarding for Satellite Appendage Vibration Control'
doc.core_properties.author='Mustafa Tolga Yavuz; Çağlar Uyulan; Sercan Acarer'
doc.core_properties.subject='Thermoelastic satellite dynamics and intelligent optimal control'

transform=etree.XSLT(etree.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
equations=[];text_log=[]
def math_element(tex,display=False):
    mm=etree.fromstring(convert(tex,display='block' if display else 'inline').encode())
    mathml={'mm':'http://www.w3.org/1998/Math/MathML'}
    for node in mm.xpath('.//mm:mover',namespaces=mathml):
        if len(node)==2 and node[1].tag.endswith('mo') and node[1].text in ['˙','¨','^','ˆ','¯','→']:
            node.set('accent','true')
    result=transform(mm);el=etree.fromstring(etree.tostring(result))
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    for prop in el.xpath('.//m:dPr',namespaces=ns):
        grow=etree.SubElement(prop,qn('m:grow'));grow.set(qn('m:val'),'1')
    for run in el.xpath('.//m:r',namespaces=ns):
        pr=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts')
        for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Cambria Math')
        pr.append(fonts);run.insert(0,pr)
    if el.tag.endswith('oMathPara') and not display:el=el.find('{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')
    return el
def add_text(p,text):
    parts=re.split(r'(\$[^$]+\$)',text)
    for part in parts:
        if part.startswith('$') and part.endswith('$'):p._p.append(math_element(part[1:-1]))
        else:p.add_run(part)
def para(text,bold=False):
    text=text.replace('37-mode',f'{eval_n}-mode').replace('37 modes',f'{eval_n} modes').replace('37 elastic',f'{eval_n} elastic').replace('37 in',f'{eval_n} in').replace('37 mode',f'{eval_n} mode').replace('76 states',f'{2*(eval_n+1)} states').replace('0.00015625 s interval',f'{eval_dt:.8f} s interval').replace('0.00015625 s step',f'{eval_dt:.8f} s step')
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;add_text(p,text)
    if bold:
        for r in p.runs:r.bold=True
    text_log.append(text);return p
def heading(text,level=1):doc.add_heading(text,level);text_log.append(text)
def eq(tex):
    number=len(equations)+1;equations.append({'number':number,'latex':tex})
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(7);p.paragraph_format.space_before=Pt(3)
    el=math_element(tex,True);p._p.append(el);p.add_run(f'   ({number})').font.size=Pt(10)
    return number
def figure(name,caption,width=6.30):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4)
    r=p.add_run();r.add_picture(str(FIG/(name+'.png')),width=Inches(width))
    for dp in p._p.xpath('.//wp:docPr'):dp.set('descr',caption)
    cap=doc.add_paragraph(style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    label,body=caption.split('. ',1);cap.add_run(label+'. ').bold=True;cap.add_run(body);text_log.append(caption)
def table(caption,headers,rows,widths):
    caption=caption.replace('37 mode',f'{eval_n} mode')
    p=doc.add_paragraph(caption,'Caption');p.paragraph_format.keep_with_next=True;text_log.append(caption)
    tb=doc.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    for i,col in enumerate(tb.columns):col.width=Cm(widths[i])
    for i,h in enumerate(headers):tb.rows[0].cells[i].text=h
    for row in rows:
        cells=tb.add_row().cells
        for c,s in zip(cells,row):c.text=str(s)
    for row in tb.rows:
        for i,c in enumerate(row.cells):
            c.width=Cm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for tag,val in [('top','80'),('bottom','80'),('left','85'),('right','85')]:
                vmar=OxmlElement('w:'+tag);vmar.set(qn('w:w'),val);vmar.set(qn('w:type'),'dxa');marg.append(vmar)
            cp.append(marg)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2);pp.paragraph_format.line_spacing=1.;pp.paragraph_format.keep_with_next=False
                for rr in pp.runs:rr.font.size=Pt(9)
    prop=tb._tbl.tblPr;bord=OxmlElement('w:tblBorders')
    for tag in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+tag);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');bord.append(x)
    prop.append(bord);repeat=OxmlElement('w:tblHeader');tb.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in tb.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for i,c in enumerate(tb.rows[0].cells):
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');c._tc.get_or_add_tcPr().append(shade)
        for pp in c.paragraphs:pp.paragraph_format.keep_with_next=True
        for r in c.paragraphs[0].runs:r.bold=True
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    text_log.extend(' | '.join(map(str,row)) for row in rows)

p=doc.add_paragraph('Thermal Preview Learning and Lyapunov Safeguarding for Satellite Appendage Vibration Control','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('Mustafa Tolga Yavuz¹   Çağlar Uyulan²   Sercan Acarer²');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('¹ Turkish Aerospace Industries, Ankara, Türkiye\n² İzmir Kâtip Çelebi University, İzmir, Türkiye');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in p.runs:r.font.size=Pt(10)
heading('Abstract')
para(fr'Eclipse heating excites flexible satellite appendages through the rate and acceleration of their thermal equilibrium. We integrate classical linear quadratic preview control with nonlinear radiative forecasting, learned feedforward compression, a conditional nominal Lyapunov filter and residual strain-rate damping. Exact optimality applies to the specified nominal unconstrained model; bounded high-order performance is assessed numerically. With five synthesis and 37 evaluation elastic modes, the common 120 s quadratic cost decreases by {100*(1-extra["nominal"]["FilteredNN_exit"]["cost"]/extra["nominal"]["LQR_exit"]["cost"]):.1f}% at exit and {100*(1-extra["nominal"]["FilteredNN_entry"]["cost"]/extra["nominal"]["LQR_entry"]["cost"]):.1f}% at entry relative to thermal LQR. In {lhs["N"]} paired Latin-hypercube cases, the median neural-to-LQR cost ratio is {lhs["paired_summary"]["FilteredNN"]["cost"]["median_ratio"]:.3f}, with paired-bootstrap 95% interval [{lhs["paired_summary"]["FilteredNN"]["cost"]["median_ratio_bootstrap95"][0]:.3f}, {lhs["paired_summary"]["FilteredNN"]["cost"]["median_ratio_bootstrap95"][1]:.3f}] under declared distributions. Horizon, network and component ablations, polynomial/RBF comparators and warmed CPU timings quantify accuracy and computation tradeoffs. Independent Hermite finite elements and a 48-cell two-face heat model expose spatial-reference mismatch. A field-aware reference with analytical residual correction reduces dynamic tip RMS from {fe["spatial_thermal"]["48"]["dynamic"]["FilteredNN"]["tip_rms_m"]*1e3:.2f} mm to {fe["spatial_thermal"]["48"]["dynamic"]["FieldNN"]["tip_rms_m"]*1e6:.2f} µm. Observer, temperature-bias and residual-delay tests reveal substantial deployment sensitivity. The 65-mode model is sampled-unstable at 6400 updates/s and requires finer intervals. The contribution is the tested thermoelastic integration and its conditional limits; no general nonlinear spacecraft safety or flight-validation claim is made.')

para(r'Keywords: thermally induced vibration; flexible satellite; piezoelectric actuation; neural optimal control; thermal preview; Lyapunov stability; eclipse transition.')

heading('1 Introduction')
para(r'Thermal transients in deployable solar arrays can disturb spacecraft attitude even when the external heating itself varies slowly compared with an elastic oscillation. Temperature differences across the panel create bending moments, and the resulting accelerations exchange angular momentum with the rigid hub. Thornton [1] describes the structural mechanisms and their spacecraft context. Johnston and Thornton [2,3] distinguish the evolving quasistatic deformation from the oscillatory component and show why derivatives of the thermal shape determine thermal-snap disturbances. This distinction identifies a control target that is physically different from maintaining a perfectly undeformed panel.')
para(r'Solar-array heating also depends on what fraction of incident radiation is converted to electricity. Xu et al. [4] explicitly include photovoltaic conversion and thermal inertia in a direct thermal-structural model. Their reported vibration reductions arise from changing the heat input, rather than from feedback control. Jiang et al. [5] examine rigid-flexible-thermal dynamics over several orbital regimes and show that drive-assembly stiffness and translational coupling can affect disturbance predictions. These studies motivate a thermal-state-aware control strategy, while also showing why a compact benchmark must state which couplings it retains.')
para(r'Piezoelectric patches provide distributed bending actuation without introducing the moving components of a separate mechanical damper. Their input matrices must nevertheless account for electrode polarity, patch geometry, added mass and stiffness, and the spatial curvature of each mode. Crawley and de Luis [6] establish the mechanical basis of this conversion. Positive position feedback [7], spacecraft smart-material control [8], adaptive-robust maneuver control [9], and passivity-based thermoelastic control [10] demonstrate different ways of combining flexible-body dynamics with actuator commands. Distributed LQR output feedback [11] addresses a further constraint: how control information can be exchanged over a prescribed communication topology.')
para(r'Intelligent optimal control offers another way to reduce online computation. Becerikli et al. [12] use dynamic neural networks to learn and prime optimized trajectories through adjoint-based training. QRnet [13] embeds an LQR contribution in a neural representation of optimal regulation, and subsequent architectures address local stability explicitly [14]. These results support retaining an analytic feedback backbone, but an accurate neural fit alone is insufficient to certify an arbitrary spacecraft controller. The public QRnet software [15] provides a useful reference for separating optimal-data generation, learning, and closed-loop evaluation; the simulations developed here use an independent implementation.')
para(r'The present study develops the thermal-vibration problem outlined in the authors’ UHUK 2026 abstract [16] into a fully specified numerical investigation. The central question is whether the nonlinear thermal forecast can be compressed into a learned optimal feedforward map while preserving a transparent LQR feedback law. The contribution lies in the application-specific integration: a moving thermal equilibrium, an analytic finite-horizon preview label, an equilibrium-vanishing neural correction, and a feasible-command safeguard. Each element is linked to an explicit equation and tested against a corresponding numerical check. No claim of priority over every thermal-control architecture is made.')
para(r'Finite preview control already separated feedback and future-information feedforward in the continuous problem of Tomizuka [18] and the discrete formulation of Tomizuka and Whitney [19]. Birla and Swarup [20] review this mature theory. The Riccati and affine-costate derivation below specializes that framework to a moving thermoelastic reference and a prescribed nonlinear thermal forecast. Its derivation is included for reproducibility, not as a claim that linear quadratic preview control is new.')
para(r'Modern command filters address learning-dependent intervention with different guarantees. Wabersich and Zeilinger [21] formulate a predictive filter using constrained model forecasts. Ames et al. [22] distinguish Lyapunov performance conditions from control barrier constraints, while Hsu et al. [23] survey the common modular structure. Our one-step box-and-half-space construction enforces a nominal Lyapunov inequality when feasible. It does not impose a state-safe invariant set or guarantee thermal, strain or attitude limits for an uncertain spacecraft.')
table('Table 1. Position of this integration within the literature',['Work','Established contribution','Difference in this study'],[
('Tomizuka [18,19] and review [20]','Classical finite preview optimal control','Nonlinear thermal forecast and moving thermoelastic coordinates'),
('Johnston and Thornton [2,3]','Thermal shape and attitude dynamics','Common-objective preview actuation and learning'),
('Meyer et al. [8]; Azadi et al. [9]; Fazelzadeh and Azadi [10]','Smart-material, adaptive-robust and thermoelastic control','Forecast-compression comparison and explicit conditional filtering'),
('Xu et al. [11]','Distributed LQR output feedback with panel experiments','Prescribed thermal transients; simulated output feedback, no hardware validation'),
('QRnet [13,14]','LQR-based neural optimal regulation and local stability','Learn only the exogenous thermal-preview correction'),
('Safety filtering [21-23]','Predictive filters and Lyapunov/barrier constructions','A nominal dissipation half-space; no barrier safety claim')],[4.0,5.7,6.9])
heading('1.1 Scope and comparison logic',2)
para(r'We consider small planar hub rotations and bending of one representative appendage. The synthesis benchmark uses two span-uniform surface temperatures. Independent validation adds 12, 24 and 48 cells per face with nonuniform heating and longitudinal conduction. The mechanical model includes hub-flexible inertial coupling and localized patch mass and stiffness. Thermal radiation remains nonlinear. The nominal incidence angle is fixed, so deformation-dependent heating, three-axis slewing, resolved eclipse geometry, hinge flexibility, and torsion remain outside the model. A prescribed spanwise shadow sweep is tested without orbital ray tracing. The uncertain simulations perturb selected physical parameters without turning this benchmark into a full nonlinear flight model.')
para(r'The principal comparison uses identical thermal coordinates, state weights, input weights, and actuator bounds for thermal LQR and neural preview. An online preview oracle measures the cost of compressing the optimal map. Frozen-forcing feedforward and polynomial/RBF surrogates assess the value of forecasting and its representation. A hub-only proportional-derivative controller provides a reference for the attitude-vibration tradeoff. It can have a smaller tip RMS while permitting a much larger hub error, so tip motion alone is not a sufficient ranking metric.')
figure('fig01_architecture','Figure 1. Physical model and control information flow. Geometry and patches follow the benchmark; deflection is exaggerated in this engineering schematic. The thermal reference precedes mechanical feedback. The learned term approximates the nominal preview map and the safeguard acts on the complete bounded command.',width=6.3)

heading('2 Thermal and mechanical model')
heading('2.1 Two surface thermal states',2)
para(r'Let the front and back temperatures be $T_f$ and $T_b$, their difference be $\Delta T=T_f-T_b$, and the effective heat capacity of each surface node per unit panel area be $C_A$. A smooth transition avoids imposing a discontinuous reference velocity. The transition direction is $d=+1$ for eclipse exit and $d=-1$ for entry, with transition center $t_e$ and width $w_e$.')
eq(r'\chi(t)=\frac{1+d\tanh((t-t_e)/w_e)}{2}')
para(r'The absorbed solar input available for heat is the incident flux multiplied by solar absorptivity minus photovoltaic conversion efficiency. Here efficiency is defined relative to the same incident optical power used by the absorptivity; this convention prevents subtracting the photovoltaic fraction twice. The nominal incidence is $\beta$, the flux multiplier is $\rho_S$, and each face receives a specified background infrared input $q_{IR}$.')
eq(r'C_A\dot T_f=(\alpha_s-\eta_{PV})S_0\rho_S\chi\cos\beta+q_{IR}-G_A(T_f-T_b)-\epsilon\sigma(T_f^4-T_{sp}^4)')
eq(r'C_A\dot T_b=q_{IR}+G_A(T_f-T_b)-\epsilon\sigma(T_b^4-T_{sp}^4)')
para(r'The conductance $G_A$ couples the nodes, and $T_{sp}=3\,\mathrm{K}$ represents the radiative sink. The Stefan-Boltzmann constant uses the CODATA value [17]. The background input is a benchmark assumption rather than an orbit-specific albedo or view-factor calculation. Summing the two balances cancels conduction exactly, which prevents internal heat transfer from becoming an artificial source of total thermal energy.')
eq(r'C_A(\dot T_f+\dot T_b)=(\alpha_s-\eta_{PV})S_0\rho_S\chi\cos\beta+2q_{IR}-\epsilon\sigma(T_f^4+T_b^4-2T_{sp}^4)')
para(r'Write the two balances as $\dot T=F_T(T,t)$. The acceleration of the thermal gradient is evaluated analytically, using the thermal Jacobian and the explicit derivative of solar incidence, rather than by differencing a noisy temperature signal.')
eq(r'\ddot T=\frac{\partial F_T}{\partial T}\dot T+\frac{\partial F_T}{\partial t},\qquad \ddot{\Delta T}=[1\quad-1]\ddot T')
para(r'The capacity is $C_A=\mu c/(2b)$, with spanwise mass density $\mu$, width $b$, and specific heat $c$. The capacity therefore corresponds to the chosen areal mass, rather than to a fictitious solid panel of facesheet-separation thickness. Symmetric actuator layers contribute to the mechanical mass and stiffness; their separate thermal capacitance, adhesive layers, and thermal-expansion mismatch are neglected in this benchmark.')

heading('2.2 Galerkin mechanical equations and piezoelectric actuation',2)
para(r'The transverse deformation is expanded in tip-normalized clamped-free beam functions. These functions provide a convenient Galerkin basis; once patches are added they need not remain exact eigenfunctions. Let $z=[\theta,q_1,\ldots,q_n]^T$, with hub angle $\theta$ and elastic coordinates $q_i$. With the adopted sign convention, positive front-minus-back temperature produces positive curvature.')
eq(r'w(x,t)=\sum_{i=1}^{n}\phi_i(x)q_i(t),\qquad \phi_i(L)=1')
eq(r'M\ddot z+C\dot z+K_m z=F_u u+g_T\Delta T')
para(r'The hub is free of an artificial torsional spring. The stiffness block for elastic coordinates is positive definite, while the complete stiffness matrix has a rigid-body null direction. Define the augmented velocity shape $\psi(x)=[R_h+x,\phi_1(x),\ldots,\phi_n(x)]^T$ and the curvature shape $\kappa(x)=[0,\phi_1^{\prime\prime}(x),\ldots,\phi_n^{\prime\prime}(x)]^T$.')
eq(r'M=J_h e_0e_0^T+\int_0^L\mu(x)\psi(x)\psi(x)^T\,dx,\qquad K_m=\int_0^L D(x)\kappa(x)\kappa(x)^T\,dx')
eq(r'g_T=\int_0^L\frac{D_0\alpha_{CTE}}{h}\kappa(x)\,dx')
para(r'The spanwise functions $\mu(x)$ and $D(x)$ include patch additions over each bonded interval. In the thermal forcing, $D_0$ is the bare panel rigidity and $h$ is the facesheet centroid separation. This idealized thermal-moment relation assumes an effective symmetric panel and ignores the actuator layers’ own differential expansion. The rigid coordinate receives no direct internal thermal torque; its disturbance arises through the off-diagonal entries of the mass matrix. Changing the coordinate sign changes both thermal and actuator signs without changing the physical predictions.')
para(r'The input vector contains one hub torque and three patch-pair voltage amplitudes. Each amplitude is applied with opposite polarity to the two patches, so the stated voltage bound applies to each patch. For a symmetric pair extending from $a_j$ to $b_j$, the modal coefficient follows from virtual work. The patch moment-per-voltage coefficient uses patch modulus $E_p$, width $b_p$, transverse strain constant $d_{31}$, and centroid distance $z_p$ from the neutral axis.')
eq(r'(F_u)_{0,0}=1,\qquad (F_u)_{i,j}=\gamma_p[\phi_i^{\prime}(b_j)-\phi_i^{\prime}(a_j)],\qquad \gamma_p=2E_p b_p|d_{31}|z_p')
para(r'Electrode polarity is chosen so that $\gamma_p$ is positive in the equation above. The generalized elastic force coefficient has units of newtons per volt. The equivalent stiffness addition of a symmetric patch pair is $2E_p b_p(h_p^3/12+h_p z_p^2)$, and the line-mass addition is $2\rho_p b_p h_p$. Structural damping is a positive semidefinite elastic block with diagonal entries $2\zeta\omega_i(M_{qq})_{ii}$, where $\omega_i$ are the bare clamped-beam angular frequencies. This damping convention is part of the benchmark, not a fitted damping model for an actual spacecraft.')
eq(r'\frac{d}{dt}\left(\frac{\dot z^TM\dot z+z^TK_mz}{2}\right)=\dot z^TF_uu+\dot z^Tg_T\Delta T-\dot z^TC\dot z')
para(r'The energy balance offers a sign and passivity check. With zero thermal forcing and zero inputs, energy cannot increase. The hub equation also preserves angular momentum when no external hub torque is applied, since neither the stiffness nor damping blocks impose an external hub load. These properties were checked before comparing any controller.')

heading('2.3 Coordinates relative to thermal equilibrium',2)
para(r'A panel held at a fixed temperature difference has a nonzero static bend. We define the desired hub angle as zero and allow the panel to follow its thermal equilibrium. Let $K_{qq}$ be the elastic stiffness block and $g_q$ the elastic thermal-force block. The reference and its derivatives are')
eq(r'r_z(t)=r_q\Delta T(t),\qquad r_q=\begin{bmatrix}0\\K_{qq}^{-1}g_q\end{bmatrix}')
eq(r'x=\begin{bmatrix}z-r_z\\\dot z-\dot r_z\end{bmatrix},\qquad \dot x=A x+B u+f_T(t)')
eq(r'A=\begin{bmatrix}0&I\\-M^{-1}K_m&-M^{-1}C\end{bmatrix},\qquad B=\begin{bmatrix}0\\M^{-1}F_u\end{bmatrix}')
eq(r'f_T=F_1\dot{\Delta T}+F_2\ddot{\Delta T},\qquad F_1=\begin{bmatrix}0\\-M^{-1}Cr_q\end{bmatrix},\quad F_2=\begin{bmatrix}0\\-r_q\end{bmatrix}')
para(r'This transformation removes the static thermal moment from the controlled dynamics. The remaining forcing depends on the rate and acceleration of the thermal gradient. It therefore connects the control problem directly to the thermal-snap mechanism identified in [3]. The target does not suppress the slowly varying static bend; a separate shape-control requirement would require another reference and another actuator-authority assessment. Thermal uncertainty or reference-estimation error enters as an additional disturbance and is not silently removed by this coordinate change.')

para(r'For the controller, a static residual-flexibility correction avoids using a low-order truncation to define the equilibrium itself. The piecewise panel rigidity gives an exact quasistatic shape per unit temperature difference. Projecting that shape onto the retained clamped-free basis calibrates the five-mode reference independently of the evaluation plant order.')
eq(r'\frac{w_s(x)}{\Delta T}=\int_0^x(x-\xi)\frac{D_0\alpha_{CTE}}{hD(\xi)}\,d\xi,\qquad \bar r_i=\frac{\int_0^L\phi_i(x)w_s(x)/\Delta T\,dx}{\int_0^L\phi_i(x)^2\,dx}')
eq(r'\bar r_q=[0,\bar r_1,\ldots,\bar r_5]^T,\qquad \bar g_T=K_m\bar r_q')
para(r'The controller replaces the truncated thermal-force vector by the calibrated vector and uses this reference in the transformed forcing. The evaluation plant retains the physical Galerkin thermal force, without this substitution. The correction modifies the nominal retained reference coefficients by less than 0.5 percent in this benchmark. It is a specified reduction correction, not an online identification result. The optimality derivation below refers to this corrected nominal five-mode system.')

heading('3 Learning an optimal thermal preview map')
heading('3.1 Finite horizon problem and analytic solution',2)
para(r'For a specified thermal forecast over a horizon $H$, define a quadratic tracking cost with a terminal penalty equal to the stabilizing infinite-horizon LQR value of the unforced mechanical system. The cost penalizes normalized attitude and elastic tracking errors as well as normalized torque and voltage. It is a control-effort objective, not a measured electrical-energy model.')
eq(r'J_H=\int_t^{t+H}(x^TQx+u^TRu)\,d\tau+x(t+H)^TPx(t+H)')
eq(r'A^TP+PA-PBR^{-1}B^TP+Q=0,\qquad K_L=R^{-1}B^TP')
para(r'Assume $Q>0$, $R>0$, and a stabilizable mechanical pair. The computed positive definite solution is selected by the stable Riccati branch. Introduce $A_c=A-BK_L$. The value function of this forced linear problem is quadratic-affine in the mechanical state, with time dependence carried by the thermal forecast.')
eq(r'V_H(t,x)=x^TPx+2s(t)^Tx+c_0(t)')
eq(r'-\dot s=A_c^Ts+Pf_T,\qquad s(t+H)=0')
eq(r'-\dot c_0=2s^Tf_T-s^TBR^{-1}B^Ts,\qquad c_0(t+H)=0')
eq(r'u_H^*(t,x)=-K_Lx+a_H(t),\qquad a_H(t)=-R^{-1}B^Ts(t)')
eq(r's(t)=\int_0^H e^{A_c^T\tau}P f_T(t+\tau)\,d\tau')
para(r'Proposition 1. Under the stated assumptions and without active input constraints, the command above is the unique minimizer of the finite-horizon problem for the supplied forecast. Proof. Substitution of the quadratic-affine value into the Hamilton-Jacobi-Bellman equation gives the Riccati equation for its quadratic terms, the terminal costate equation for its linear terms, and the scalar equation for its constant term. The Hamiltonian is strictly convex in the input because $R>0$. Its stationary point is therefore the unique minimizer. The terminal conditions agree with the specified terminal penalty, completing the verification argument.')
para(r'The feedback gain is fixed; only the feedforward map depends nonlinearly on the thermal state, photovoltaic efficiency, and eclipse timing. Thus learning does not need to approximate the entire mechanical feedback law. The optimality claim applies to the forced linear mechanical model with a prescribed nonlinear thermal forecast. It does not establish optimality for a deformation-dependent nonlinear mechanical or radiation model, and clipping or filtering changes the finite-horizon optimizer.')

heading('3.2 Optimal data and neural architecture',2)
para(r'Training labels are generated by integrating the nonlinear thermal model and evaluating the costate integral. The thermal input vector contains normalized mean temperature, temperature difference, transformed time to the transition, transition direction, photovoltaic efficiency, and solar-flux multiplier. A hyperbolic tangent transformation bounds the time feature while retaining resolution near the event.')
eq(r'\xi_T=\begin{bmatrix}(T_f+T_b-560)/110\\(T_f-T_b)/40\\\tanh((t-t_e)/18)\\d\\(\eta_{PV}-0.30)/0.08\\(\rho_S-1)/0.20\end{bmatrix}')
para(fr'The network has {learn["architecture"][0]} inputs, hidden layers of {learn["architecture"][1]} and {learn["architecture"][2]} hyperbolic-tangent units, and {learn["architecture"][3]} linear outputs. Output scaling uses the training-label standard deviation. The {learn["N"]} independent thermal states are split into {learn["split"][0]} training, {learn["split"][1]} validation, and {learn["split"][2]} test states using seed {learn["seed"]}. Mean temperatures span 218 to 370 K, differences span -30 to 35 K, transition times span -55 to 65 s relative to the event, efficiencies span 0.22 to 0.38, and flux multipliers span 0.85 to 1.15. These ranges define a synthetic training envelope; they are not measured operating limits.')
eq(r'\mathcal L=\frac{1}{4N_{tr}}\sum_{k=1}^{N_{tr}}\left\|D_a^{-1}\big(N_\vartheta(\xi_T^{(k)})-a_H^{(k)}\big)\right\|_2^2')
para(fr'Adam optimization runs for {learn["epochs"]} epochs, with batches of approximately 256 states, first and second moment factors 0.9 and 0.999, and a scheduled learning rate decreasing from 0.0025. The stored weights are selected by validation error. Test labels do not influence training or epoch selection. The validation split assesses interpolation within the stated envelope; physical-parameter uncertainty is assessed separately in closed-loop trials.')
para(r'A smooth activity gate makes the learned correction vanish when the thermal subsystem is exactly at a steady state with no changing illumination. Without this condition, a small neural bias can impose a persistent mechanical command even when the desired thermal forcing is zero. The gate uses the analytic thermal rates and the illumination derivative, with a small regularization scale.')
eq(r'\gamma_T=\frac{\|\dot T\|_2^2+(0.3\dot\chi)^2}{\|\dot T\|_2^2+(0.3\dot\chi)^2+10^{-8}},\qquad a_N=\gamma_TN_\vartheta(\xi_T)')
eq(r'u_N=-K_Lx+a_N')
para(r'Temperatures and their analytic rates are expressed in kelvin and seconds in the gate, and the coefficient 0.3 has units of kelvin. The threshold is therefore an explicitly specified numerical regularization, rather than a dimensionless physical law. It affects only very small thermal activity in the simulated transitions. The analytic LQR derivative with respect to the mechanical state remains unchanged because the learned map depends on thermal variables alone.')

heading('3.3 What near optimality means',2)
para(r'For a disturbance that tends to zero sufficiently fast, let $a_\infty$ denote the infinite-preview costate command. Completing the square in the verified infinite-horizon value gives an exact excess-cost identity for any admissible nominal trajectory with the required terminal decay.')
eq(r'J(u)-J(u_\infty^*)=\int_0^\infty\big(u+K_Lx-a_\infty\big)^TR\big(u+K_Lx-a_\infty\big)\,dt')
para(r'Consequently, the approximation loss of the unfiltered learned command is determined by the weighted preview error, provided the model and forecast assumptions hold. A filter adds its own intervention error. This statement separates a mathematical near-optimality measure from a low neural regression error and from an empirical ratio of closed-loop costs.')
eq(r'\|a_H-a_\infty\|\le\frac{\|R^{-1}B^T\|\,M_c\,\|P\|\,\bar f}{\lambda_c}e^{-\lambda_cH}')
para(r'The tail estimate assumes $\|e^{A_c^Tt}\|\le M_c e^{-\lambda_ct}$ and $\|f_T\|\le\bar f$. A stable spectral abscissa alone does not set the transient multiplier $M_c$, particularly for a nonnormal matrix. The reported regression errors are finite-sample statistics and are not substituted for a uniform error bound. Likewise, the receding-horizon oracle used in the experiments supplies a numerical comparator; it is not labeled as the global optimizer of the full 37-mode, saturated, uncertain plant.')

heading('4 Feasible Lyapunov safeguard')
heading('4.1 Nominal dissipation condition and blending rule',2)
para(r'Let $V_L=x^TPx$ and $W=Q+K_L^TRK_L$. The nominal LQR obeys a dissipation identity with an explicit thermal-disturbance term. We require a fraction $\sigma_L=0.25$ of its unforced dissipation while allowing the worst sign of the instantaneous thermal work. The use of a known forcing term avoids demanding an impossible monotone decrease during every thermal transient.')
eq(r'\dot V_L(u_L)=-x^TWx+2x^TPf_T')
eq(r'\dot V_L(u)\le-\sigma_Lx^TWx+2|x^TPf_T|')
eq(r'a_c=2B^TPx,\qquad b_c=-\sigma_Lx^TWx+2|x^TPf_T|-2x^TP(Ax+f_T)')
para(r'The admissible command set is the actuator box intersected with one affine inequality. Its exact feasibility condition follows by minimizing the linear form over that box.')
eq(r'\mathcal U_c=\{u:|u_i|\le\bar u_i,\ a_c^Tu\le b_c\},\qquad -\sum_i\bar u_i|a_{c,i}|\le b_c')
para(r'First clip the learned command and the LQR baseline to the actuator limits. If the clipped learned command is admissible, apply it. When the baseline is admissible and the learned command is not, restrict the intervention to the segment between them. The largest admissible blend is the closest point to the learned command on that segment in any positive definite quadratic input metric.')
eq(r'u_F=u_b+\alpha^*(u_n-u_b),\qquad \alpha^*=\min\left(1,\frac{b_c-a_c^Tu_b}{a_c^T(u_n-u_b)}\right)')
para(r'This formula is used only in the violating-candidate, feasible-baseline case, so its denominator is positive and the blend lies between zero and one. If clipping makes the baseline infeasible, the implementation instead solves the box-constrained weighted projection onto the half-space by a scalar dual search. If the intersection is empty, it returns the box command minimizing the inequality’s left side and records the remaining deficit. Such a step is a declared loss of the certificate, not a guaranteed stable fallback. No infeasible intersections occurred in the reported nominal or uncertainty runs.')
para(r'A segment intervention limits unnecessary redistribution between torque and piezoelectric channels. It also prevents the safeguard from introducing a larger corrective command than the two endpoints when the nominal baseline remains feasible. This restriction is useful for a flexible plant because a certificate derived from a finite set of modes does not certify every neglected mode. The higher-order plant and refinement checks therefore remain essential.')

heading('4.2 Conditional stability statement',2)
para(r'Proposition 2. For the nominal continuous-time model, suppose the command satisfies the dissipation inequality at every time, the intersection remains feasible, and the chosen closed-loop solution is well posed. Then the mechanical tracking error is input-to-state stable with respect to the thermal forcing. If the thermal forcing tends to zero, the state converges to zero. To express quantitative bounds without mixing physical coordinate scales, the following eigenvalues and norms are evaluated after the nondimensional state transformation used for the Riccati solve.')
eq(r'\dot V_L\le-\frac{\sigma_L\lambda_{min}(W)}{2\lambda_{max}(P)}V_L+\frac{2\|P\|^2}{\sigma_L\lambda_{min}(W)}\|f_T\|^2')
para(r'Proof. Bound the thermal work by $2\|P\|\|x\|\|f_T\|$ and apply Young’s inequality with half of the available dissipation. Use positive definiteness of $P$ and $W$ to obtain the scalar differential inequality above. Integration gives an exponentially decaying initial-state term and a convolution of the squared disturbance. The equilibrium gate removes the learned feedforward bias at a steady thermal equilibrium. The conclusion requires continuous enforcement and well-posed dynamics; finite-step numerical evaluation is supporting evidence rather than a proof of these implementation properties.')
para(r'For an additional modeling perturbation $\delta(x,t)$ satisfying a local bound $\|\delta\|\le\ell\|x\|+d_\delta(t)$, the derivative gains a term $2x^TP\delta$. A sufficient local small-gain condition is')
eq(r'2\|P\|\ell<\sigma_L\lambda_{min}(W)')
para(r'Under that condition, the same proof yields an input-to-state bound driven by $\|f_T\|+d_\delta$. The paper does not identify a rigorous global value of $\ell$ for all physical uncertainty, neglected modes, sensor errors, and saturation cases. The Monte Carlo study therefore supplies empirical robustness evidence within its stated ranges, while the analytical claim remains conditional. In particular, a nominal Lyapunov filter cannot by itself guarantee stability of a thermal-flutter model or of arbitrary unmodeled elastic modes.')

heading('4.3 Damping measured residual strain rates',2)
para(r'Increasing the number of modes used for evaluation reveals a second issue: a truncated state-feedback law can excite modes outside the observer or retained-coordinate model. A piezoelectric sensor can nevertheless measure a strain-rate combination containing those modes. The implementation therefore uses the difference between the full calibrated power-conjugate piezoelectric velocity output and its prediction from the retained coordinates. This difference is available from the sensor output without separately estimating every omitted modal coordinate.')
eq(r'y_p=F_{u,p}^T\dot z,\qquad \hat y_p=F_{u,p,5}^T\dot{\hat z},\qquad u_D=\begin{bmatrix}0\\-k_D(y_p-\hat y_p)\end{bmatrix}')
para(r'Here the subscript p selects the three patch channels and the hub-torque component is zero. The gain is 5000 V² s per N m. Sensor calibration and instantaneous strain-rate availability are assumptions of the numerical experiment. The same residual-damping channel is used with thermal LQR, the preview oracle, frozen-forcing and surrogate feedforward, and both learned controllers. The hub-only reference has no patch actuation. Under the exact five-mode nominal model the measured residual is zero, so the finite-horizon optimality result is unchanged.')
para(r'The physical candidate and baseline become the core command plus the residual-damping command. For the Lyapunov inequality, its known nominal contribution is treated as part of the effective disturbance. This preserves the feasible-baseline construction without allowing the safeguard to remove a useful high-mode damping command solely because it was absent from the low-order model.')
eq(r'f_{eff}=f_T+Bu_D,\qquad u_{b}=-K_Lx+u_D,\qquad u_{n}=-K_Lx+a_N+u_D')
eq(r'a_c^T u_{tot}\le -\sigma_Lx^TWx+2|x^TP(f_T+Bu_D)|-2x^TP(Ax+f_T)')
eq(r'a_c^T v\le -\sigma_Lx^TWx+2|x^TPf_{eff}|-2x^TP(Ax+f_{eff}),\quad v=u_{tot}-u_D')
para(r'The first inequality uses the complete physical command $u_{tot}$ inside the physical derivative. The equivalent second form uses the core command $v$ and the shifted bounds $-\bar u-u_D\le v\le\bar u-u_D$. Only the allowed disturbance work uses $f_{eff}$ in the complete-command form. Adding $Bu_D$ there to the physical derivative while also retaining the complete command would double count residual damping. An independent 10000-case algebra check compares the two residuals.')
para(r'In the bounded-command construction, the disturbance-work allowance uses the effective disturbance, while the physical derivative retains the thermal forcing and the complete input. The stability proposition therefore applies to the retained nominal coordinates with the effective disturbance. Its full-plant interpretation still needs a small-gain or passivity argument that is not established here. Continuous and sampled high-order eigenvalue checks, deliberate corruption tests, and modal refinement provide the stated numerical evidence. The residual channel is not presented as a universal certificate for an infinite-dimensional structure.')

heading('5 Numerical experiment and verification protocol')
heading('5.1 Source informed benchmark',2)
para(r'The panel length, width, line mass, rigidity, hub size, and solar-flux scale are informed by the illustrative solar-panel model in [3]. Other entries are explicit benchmark design assumptions, chosen to produce a finite thermal gradient and a realizable patch model. The panel is not presented as a reconstructed flight article. Thermal optical properties and specific heat are informed by [10], and photovoltaic efficiency follows the modeling motivation in [4]. Table 2 distinguishes those origins.')
table('Table 2. Nominal physical and actuator parameters',['Quantity','Value and units','Basis'],[
('Panel length and width','9 m and 3 m','Illustrative geometry in [3]'),('Bare line mass and rigidity','7.4 kg m⁻¹ and 2.0 × 10⁴ N m²','Illustrative model in [3]'),('Hub radius and inertia','1 m and 2500 kg m²','Radius in [3]; inertia of a 5000 kg solid cylinder assumption'),('Facesheet separation and effective CTE','0.025 m and 2.0 × 10⁻⁶ K⁻¹','Design assumptions'),('Specific heat and node capacity','1044 J kg⁻¹ K⁻¹ and 1287.6 J m⁻² K⁻¹','Specific heat informed by [10]; capacity from chosen areal mass'),('Solar absorptivity and emissivity','0.92 and 0.82','Optical-property values in [10]'),('Solar flux and nominal PV efficiency','1350 W m⁻² and 0.30','Flux scale in [3]; PV motivation in [4]'),('Node conductance and background IR','35 W m⁻² K⁻¹ and 120 W m⁻² per face','Design assumptions'),('Incidence and transition width','20 degrees and 3 s','Design assumptions; 5 to 95 percent transition is 8.83 s'),('Structural damping ratio','0.002','Low-damping benchmark assumption'),('Patch modulus and strain constant','63 GPa and |d₃₁| = 190 pm V⁻¹','Assumed PZT properties'),('Patch density width and thickness','7800 kg m⁻³; 0.30 m; 0.30 mm','Assumed PZT properties and geometry'),('Patch intervals','[0, 0.5], [1.8, 2.3], [4.5, 5.0] m','Chosen placements'),('Hub torque and patch voltage bounds','±0.30 N m and ±150 V','Chosen input limits')],[4.1,6.3,6.2])
para(r'The controller uses five elastic basis functions plus one rigid coordinate, producing a twelve-state mechanical model. The evaluation plant uses 37 elastic functions plus the rigid coordinate, producing 76 states. The patch contributions are integrated by 60-point Gaussian quadrature for synthesis and 100-point quadrature for high-order evaluation on subintervals split at every patch edge. The lowest five patched clamped frequencies are '+', '.join(f'{f:.4f}' for f in v['clamped_patched_frequencies_hz'])+' Hz. They are clamped-panel frequencies; they are not mislabeled as the eigenfrequencies of the freely rotating complete spacecraft.')
eq(r'D_x=\mathrm{diag}(D_z,D_v),\qquad D_z=\mathrm{diag}(10^{-4},0.002,0.002,0.002,0.02,0.02)')
eq(r'D_v=\mathrm{diag}(5\times10^{-4},0.005,0.005,0.005,0.05,0.05)')
eq(r'Q=D_x^{-T}D_x^{-1},\qquad R=\mathrm{diag}(0.30,150,150,150)^{-2}')
para(r'The state ordering is hub angle, five elastic displacements, hub angular rate, and five elastic velocities. The last two elastic coordinates receive lower tracking weights so that the objective emphasizes the three lower modes while still retaining high-mode dynamics in the synthesis. Riccati computations are performed in normalized coordinates and transformed back to SI units. This step avoids misleading numerical conditioning from combining radians, meters, torque, and voltage in one matrix calculation.')

heading('5.2 Comparative experiments and paired uncertainty',2)
para(r'The revised comparisons use a common 120 s interval including initialization, with eclipse exit and entry centered at 20 s. RMS and command moments use every native sample from 0 through 120 s; the running-cost rectangle rule excludes the terminal sample. This differs from the initial study’s 180 s interval and 5 s RMS exclusion. Temperatures begin at the dark or sunlit nonlinear equilibrium and mechanical coordinates begin at the corresponding static shape and rate. No unrelated mechanical kick is introduced.')
para(r'The oracle forecasts the two-node nonlinear thermal model for horizons 6, 12, 24, 36 and 48 s, with 0.12 s midpoint quadrature and exactly integrated mechanical kernels. The zero-horizon limit is thermal LQR. The oracle grid is 0.1 s; current temperatures on that grid come from the actual thermal plant, whereas all future predictions retain nominal heat-transfer parameters. The frozen-forcing comparator replaces the future forcing by its instantaneous value throughout the 24 s integral, thereby separating future information from a static feedforward correction. Polynomial regression uses all monomials of total degree at most three in the same six features and validation-selected ridge regularization.')
para(r'A stronger non-neural comparator uses 256 Gaussian radial basis centers selected by 30 k-means++ iterations on training features only, with seed 20261005. The basis uses $\xi=\xi_T$ and includes an intercept and the six linear features. Width and ridge values are selected on the unchanged validation set from six inverse-width settings and three ridge settings; the test set remains unused in selection. This comparison supplements the inexpensive cubic polynomial with a coefficient budget closer to the neural map.')
eq(r'a_{RBF}(\xi)=\begin{bmatrix}1&\xi^T&e^{-\beta_R\|\xi-c_1\|^2}&\cdots&e^{-\beta_R\|\xi-c_{256}\|^2}\end{bmatrix}\Theta')
eq(r'a_{fr}=-R^{-1}B^T\left[\int_0^H e^{A_c^T\tau}\,d\tau\right]Pf_T(t)')
para(r'The passivity comparator combines the same hub PD gains, 150 N m per radian and 550 N m s per radian, with collocated patch velocity damping $u_p=-5000F_{u,p}^T\dot z$. Under ideal continuous sensing, the patch contribution to structural energy is $-5000\|F_{u,p}^T\dot z\|^2$. Its objective and feasible authority are assessed with the same reported cost; it is a damped reference rather than a preview optimizer. The label Hub PD replaces the ambiguous label Passive used previously.')
para(fr'The uncertainty design contains {lhs["N"]} Latin-hypercube points with seed {lhs["seed"]}, including a balanced binary entry/exit coordinate. Every point is evaluated under thermal LQR, 24 s oracle, raw gated NN and filtered gated NN, with common plant parameters, initial temperatures and actuator bounds. Uniform marginal ranges are declared in Table 3; they are assumptions for a sensitivity benchmark, not identified population distributions. A paired bootstrap resamples complete case indices 10000 times using seed 55041. The confidence intervals quantify sampling variability of a median ratio or mean paired difference within this design, not spacecraft reliability.')
table('Table 3. Latin hypercube uncertainty design',['Quantity','Uniform range'],[('Rigidity D / D0','0.8 to 1.2'),('Line mass mu / mu0','0.9 to 1.1'),('Specific heat c / c0','0.8 to 1.2'),('Conductance G / G0','0.7 to 1.3'),('Damping zeta / zeta0','0.5 to 1.5'),('Piezoelectric actuation effectiveness','0.7 to 1.0'),('Photovoltaic efficiency','0.22 to 0.38'),('Solar irradiance multiplier','0.9 to 1.1'),('Eclipse direction','Equal LHS strata mapped to entry and exit')],[8,8.6])
para(r'Two additional correlated scenarios change effective thickness by factors 0.85 and 1.15, with rigidity proportional to thickness cubed, line mass proportional to thickness, transverse conductance inversely proportional to thickness, and unchanged specific heat. Patch lever arms are recomputed. Two corners combine extreme rigidity, mass, thermal capacity and conductance with half nominal damping and 70 percent actuation effectiveness. These idealized correlations test a physically interpretable perturbation; they do not replace laminate characterization.')

heading('5.3 Observer and implementation perturbations',2)
para(r'Output feedback estimates the physical retained displacement and velocity coordinates. Measurements consist of hub angle, angular rate and three patch-averaged bending strains. The true sensor map includes every plant mode; the estimator uses the nominal five-mode map. A dual discrete Riccati design is solved in normalized coordinates, with process covariance $10^{-6}I$ per 1 ms innovation update and the low-noise covariance in Table 4. Between updates the estimator propagates the nominal physical model with the applied input and measured thermal difference. It is initialized at the true equilibrium to isolate noise, bias and subsequent model/sensor spillover; initial estimation uncertainty remains untested.')
eq(r'\hat y_{k}^{+}=\hat y_{k}^{-}+L_o(m_k-C_o\hat y_{k}^{-}),\qquad \hat y_{j+1}^{-}=A_{o,d}\hat y_j^{+}+B_{o,d}u_j+E_{o,d}\bar\Delta_j')
para(r'Independent Gaussian draws are held for 1 ms; the same seed and biases are used for each controller pair. Analytic thermal derivatives are evaluated from measured temperatures and the nominal heat model, avoiding differentiation of white noise. The safeguard uses estimated states; its nominal expression evaluated at true retained states is reported separately. Neither expression is the exact derivative of the high-order uncertain plant. The observer study preserves an ideal residual-rate channel, followed by a separate study adding noise and persistent bias to that calibrated channel.')
table('Table 4. Assumed sensing levels and implementation tests',['Quantity','Low case','High case'],[('Surface temperature bias per face','+0.1 / -0.1 K','+0.5 / -0.5 K'),('Temperature noise standard deviation','0.05 K','0.2 K'),('Angle and gyro noise standard deviations','0.02 arcsec; 0.02 arcsec/s','0.1 arcsec; 0.1 arcsec/s'),('Angle and gyro biases','+0.02 arcsec; -0.02 arcsec/s','+0.1 arcsec; -0.1 arcsec/s'),('Patch strain noise and bias magnitude','5 nanostrain','25 nanostrain'),('Separate residual port noise standard deviation','1e-6 N m / (V s)','1e-5 N m / (V s)'),('Residual port bias','Half standard deviation, fixed signs','Half standard deviation, fixed signs'),('Core command delay','0, 3, 6, 13 native intervals','0 to 2.03125 ms'),('Residual command delay','1, 2, 3 native intervals','0.15625 to 0.46875 ms'),('First order voltage and torque driver','50, 200 and 1000 Hz bandwidth','Same bounds and filter')],[6.6,5,5])
para(r'Delays are modeled by command buffers initialized at zero, separately for the core feedback/feedforward command and the residual damping command. Delayed candidates are bounded at the current evaluation instant. A first-order driver follows the resulting command with discrete coefficient $1-\exp(-2\pi f_b\Delta t)$. The filter acts before this driver, so driver lag can invalidate its instantaneous applied-command inequality. Numerical termination is recorded if cost density exceeds $10^{12}$ or tip deviation exceeds 0.5 m. Bounded finite trajectories are not automatically classified as stable; sampled linear eigenvalue violations are reported independently.')

heading('5.4 Independent structural and spatial thermal models',2)
para(r'The independent mechanical model assembles cubic Hermite displacement/slope elements with consistent mass, local bare-plus-patch rigidity and mass, exact hub-flexible inertia integrals, thermal moments and patch curvature integrals. Patch edges coincide with mesh boundaries. Five-point Gaussian quadrature integrates each element, split again at thermal-cell boundaries. Shift-invert generalized eigensolution extracts 37 patched clamped modes; these coordinates are distinct from the analytic cantilever trial basis. A weighted spatial projection maps the FE displacement field to the five synthesis coordinates. This independently validates discretization and coupling under the same planar Euler-Bernoulli physics; it is not a laboratory or flight validation.')
eq(r'N_e(\zeta)=\begin{bmatrix}1-3\zeta^2+2\zeta^3&\ell_e\zeta(1-\zeta)^2&\zeta^2(3-2\zeta)&\ell_e\zeta^2(\zeta-1)\end{bmatrix}')
eq(r'M_e=\int_e\mu_eN_e^TN_e\,dx,\qquad K_e=\int_eD_e(N_e^{\prime\prime})^TN_e^{\prime\prime}\,dx')
para(r'The FE model uses modal viscous damping $2\zeta\omega_j$ in mass-normalized clamped coordinates, whereas the analytic trial-basis model uses a diagonal bare-frequency damping prescription. This distinction is intentional and disclosed. Meshes requested at 60, 120, 240, 480 and 960 subdivisions gain a few nodes at patch boundaries. At very fine meshes a dense lowest-eigenvalue extraction loses relative accuracy; all final FE results use shift-invert extraction. Bare frequencies, exact piecewise static bending, unforced energy decay and hub angular momentum supply independent checks.')
para(r'A conservative two-face finite-volume thermal extension uses 12, 24 and 48 uniform cells per face. Longitudinal coefficient $k_\parallel h=0.5$ W/K is an assumed effective value. The heat-flux multiplier varies as $1+0.25(2x/L-1)$ and a prescribed shadow sweep delays the local transition by $3x/L$ seconds. End faces have zero longitudinal heat flux. This constructed spatial transient is an exploratory validation case, not an imported spacecraft temperature field.')
eq(r'q_{\parallel,f,i}=\frac{k_\parallel h}{\Delta x^2}(T_{f,i-1}-2T_{f,i}+T_{f,i+1})')
eq(r'C_A\dot T_{f,i}=q_{s,i}+q_{IR}-G(T_{f,i}-T_{b,i})-\epsilon\sigma(T_{f,i}^4-T_s^4)+q_{\parallel,f,i}')
para(r'The back-face equation has no direct solar term and the opposite transverse-conduction term. At each insulated end the missing neighbor is replaced by the boundary-cell temperature. Longitudinal and transverse conduction cancel in the summed energy balance. Thermal trajectories use an independent adaptive integrator and are interpolated on a 0.01 s grid; mechanical propagation retains the native interval and the midpoint thermal moment for every cell.')
para(r'We first expose the error from supplying only mean surface temperatures to the original reference. A second integration uses the projected spatial static shape and its rate as the retained reference. It combines the original scalar neural map with a frozen-forcing residual correction, rather than claiming that the six-input network learned arbitrary thermal fields. The residual forcing is the difference between spatial reference forcing and the mean-field forcing. Its added input is obtained from the same integrated mechanical kernel. This strategy requires field-mode temperature information and a structural reference map; three sensor types alone do not establish its field observability.')
eq(r'r_s=H_q\begin{bmatrix}0\\K_{FE}^{-1}g_{FE}(\Delta_1,\ldots,\Delta_n)\end{bmatrix},\quad f_s=\begin{bmatrix}0\\-M^{-1}C\dot r_s-\ddot r_s\end{bmatrix}')
eq(r'a_{mix}=a_N-R^{-1}B^T\left[\int_0^H e^{A_c^T\tau}\,d\tau\right]P(f_s-f_{mean})')
para(r'Spatial reference rates come from the heat-model right-hand side. Reference acceleration is a second-order central difference of those smooth rates on the 0.01 s grid, with one-sided endpoint formulas. This numerical approximation does not extend the exact optimality proposition to the FE plant. Both field-aware controllers use the same spatial reference, weights, bounds and residual damping.')
para(r'The integration can also be understood as a structured approximation to a spatial preview teacher. For the same nominal linear mechanical model, with thermal force chosen consistently with the spatial reference and with a prescribed spatial heat forecast, the HJB derivation remains valid after replacing $f_T$ by $f_s$. Its preview map splits linearly into mean and residual forcing contributions. The implemented mixed command learns the first part and freezes the current residual for the second; it is therefore a specified approximation to this teacher, not a new general optimal-control law.')
eq(r'a_H^s=a_H^{mean}-R^{-1}B^T\int_0^H e^{A_c^T\tau}P f_r(t+\tau)\,d\tau,\qquad f_r=f_s-f_{mean}')
para(r'If the residual forcing satisfies $\|f_r(t+\tau)-f_r(t)\|\le L_r\tau$ and the same exponential bound used in Section 3.3 holds, triangle inequality yields a pointwise approximation bound. The learned mean-map error includes the activity gate. This explains why a slowly varying spatial residual can be handled cheaply, while a rapidly sweeping shadow may require an actual field forecast.')
eq(r'\|a_{mix}-a_H^s\|\le\|a_N-a_H^{mean}\|+\|R^{-1}B^T\|\,\|P\|M_cL_r I_1(H)')
eq(r'I_1(H)=\frac{1-(1+\lambda_cH)e^{-\lambda_cH}}{\lambda_c^2}')
para(r'This bound follows by integrating $\tau e^{-\lambda_c\tau}$. It is conditional on a residual-rate bound and applies to the nominal teacher map. The present simulations do not identify a uniform $L_r$ or claim a certified mixed-controller cost gap on the FE plant. They test the proposed decomposition on one stated spatial transient and retain the scalar-reference failure alongside it.')

heading('5.5 Mathematical and numerical checks',2)
table('Table 5. Independent checks and interpretation',['Check','Value','Interpretation'],[
('Nominal CARE relative residual',f'{v["care_relative_residual"]:.3e}','Scaled solve, physical residual'),
('Independent MATLAB P relative difference',f'{mv["P_relative_difference"]:.3e}','Independent CARE solver'),
('Independent MATLAB trajectory scaled maximum difference',f'{mv["trajectory_max_scaled_difference"]:.3e}','ode113 versus DOP853'),
('Independent HJB identity maximum relative residual',f'{mv["HJB_identity_max_relative_residual"]:.3e}','Nominal quadratic affine identity'),
('Cached versus original analytic preview maximum absolute difference',f'{json.loads((DATA/"streaming_validation.json").read_text())["cached_preview_max_abs"]:.3e}','Three thermal states, all channels'),
('New versus previous native trace maximum difference',f'{max(json.loads((DATA/"streaming_validation.json").read_text())[k]["native_trace_max_abs"] for k in ["LQR","RawNN","FilteredNN"]):.3e}','30 s thermal transition, three control laws'),
('Complete versus core command half space maximum difference',f'{refine["algebra"]["max_residual_difference"]:.3e}','10000 independent random tests'),
('Independent FE maximum unforced energy increment',f'{refine["FE_conservation"]["maximum_energy_increase_J"]:.3e} J','Nonincreasing structural energy'),
('Independent FE maximum hub momentum drift',f'{refine["FE_conservation"]["momentum_max_abs_drift"]:.3e}','Zero external torque'),
('Spatial heat maximum summed balance error',f'{max(abs(vv["energy_balance_error_W"]) for vv in fe["spatial_thermal"].values()):.3e} W','Conduction cancellation and surface input')],[7.2,3,6.4])
para(r'The MATLAB check independently solves algebra and trajectories from exported nominal matrices; the FE assembly provides the separate structural discretization check. Earlier time halving changed tip RMS by approximately $3.1\times10^{-5}$ relatively on the 37-mode model. The final modal study extends through 65 modes and explicitly reevaluates sample stability. These checks address distinct errors and must not be treated as one global model-validation certificate.')

heading('6 Results and interpretation')
heading('6.1 Common objective comparators and preview horizon',2)
figure('fig02_thermal','Figure 2. Nonlinear face temperatures and through-thickness difference during entry and exit. The same prescribed thermal transition supplies every nominal comparator.')
rows=[]
for suffix,label in [('exit','Exit'),('entry','Entry')]:
    for name in ['HubPD','Passivity','LQR','Frozen24','Oracle24','Polynomial','RBF256','RawNN','FilteredNN']:
        a=rbf_report['nominal'][suffix] if name=='RBF256' else extra['nominal'][name+'_'+suffix];rows.append((label,name,f'{a["tip_rms_m"]*1e6:.3f}',f'{a["angle_rms_arcsec"]:.4f}',f'{a["cost"]:.5g}'))
table('Table 6. Nominal 120 s results with identical reported weights and bounds',['Transition','Controller','Tip RMS µm','Angle RMS arcsec','Cost'],rows,[2,3.3,3.7,4.2,3.4])
base=extra['nominal']['LQR_exit'];neural=extra['nominal']['FilteredNN_exit'];oracle=extra['nominal']['Oracle24_exit']
para(fr'At exit, filtered NN reduces the common cost from {base["cost"]:.5f} to {neural["cost"]:.5f}, dynamic tip RMS from {base["tip_rms_m"]*1e6:.3f} to {neural["tip_rms_m"]*1e6:.3f} µm, and hub-angle RMS from {base["angle_rms_arcsec"]:.4f} to {neural["angle_rms_arcsec"]:.4f} arcsec. Its cost is {neural["cost"]/oracle["cost"]:.4f} times the nominal 24 s oracle comparator. These are high-order sampled costs, so even a value below the oracle would not establish superiority over the optimum of the identical five-mode unconstrained problem.')
figure('fig03_exit_response','Figure 3. Exit responses for thermal LQR, frozen-forcing feedforward, analytic preview and filtered neural preview. The plotted 90 s window resolves the transient; Table 6 uses all 120 s.')
figure('fig04_entry_response','Figure 4. Entry responses using the same cost and input bounds. Initial temperatures satisfy the nonlinear sunlit balance.')
rows=[]
for H in [0,6,12,24,36,48]:
        a=extra['horizons'][f'H{H}_exit'];rows.append((H,f'{a["cost"]:.6f}',f'{a["tip_rms_m"]*1e6:.3f}',f'{a["angle_rms_arcsec"]:.4f}',f'{np.sum(np.array(a["channels"])[1]**2/np.array([.3,150,150,150])**2):.3e}'))
table('Table 7. Exit analytic preview horizon sensitivity',['H s','Cost','Tip RMS µm','Angle RMS arcsec','Weighted input mean square'],rows,[2,3.2,3.3,4.0,4.1])
para(r'The horizon study includes both transition directions and input effort. Most cost improvement already occurs at 6 s; 12 s nearly attains the cost plateau. The 24 s setting is a conservative label-generation horizon, not a uniquely optimal tuning choice. Tip RMS does not monotonically decrease with horizon because the optimized cost combines retained state error and actuator effort, while the reported tip also contains omitted-mode motion. A shorter horizon or a simpler surrogate can be preferable when computation or memory dominates.')

heading('6.2 Learning and component sensitivity',2)
para(fr'The original 6000-label network has {100*learn["test_relative_L2"]:.3f}% scaled held-out relative L2 error before the activity gate. Figure 5 shows every output, including hub torque. The cubic polynomial uses {extra["polynomial_learning"]["basis_terms"]} basis terms and has {100*extra["polynomial_learning"]["test_scaled_relative_L2"]:.2f}% scaled test error on the identical held-out points. Its closed-loop results in Table 6 quantify the consequences of this simpler approximation rather than assuming that a neural representation is necessary.')
para(fr'The 256-center RBF fit has {100*rbf_report["test_scaled_relative_L2"]:.2f}% scaled held-out error with validation-selected inverse width {rbf_report["gamma"]:.3g} and ridge {rbf_report["ridge"]:.3g}. Its exit cost is {rbf_report["nominal"]["exit"]["cost"]:.5f}. This provides a stronger basis for judging forecast compression than regression error from the cubic polynomial alone. Learning a neural map is an implementation choice whose accuracy, cost and timing must be compared with these alternatives.')
figure('fig05_learning','Figure 5. Four-channel held-out parity plots. Each panel compares analytic preview labels with predictions before gating and filtering; the diagonal denotes exact agreement. These are synthetic training targets, not experimental measurements.')
rows=[]
for a in extra['networks']:
    c=a['closed_loop'];rows.append((str(a['hidden']),a['N_training'],a['epochs'],a['activation'],'yes' if a['output_scaling'] else 'no',f'{100*a["test_scaled_relative_L2"]:.2f}',f'{c["cost"]:.5g}'))
table('Table 8. One factor at a time learning sensitivity',['Hidden units','Training N','Epochs','Activation','Output scale','Test error %','Exit cost'],rows,[2.6,2.2,1.7,2.4,2.1,2.8,2.8])
para(r'Nested training subsets contain 1050, 2100 and 4200 samples from the original fixed split; validation and test sets remain at 900 each. Reported architecture, epoch, activation and output-scaling variants share labels and evaluation points. They use one specified initialization seed, so the results describe a sensitivity study rather than a statistical neural-hyperparameter optimum. Output scaling is varied; an exhaustive input-normalization study is not claimed.')
rows=[]
for label,a in extra['ablations'].items():rows.append((label,f'{a["cost"]:.5g}',f'{a["tip_rms_m"]*1e6:.2f}',f'{a["filter_fraction"]*100:.2f}',f'{a["correction_max_normalized"]:.3g}'))
table('Table 9. Component and corruption ablations',['Case','Cost','Tip RMS µm','Filter active %','Maximum normalized correction'],rows,[3.8,3,3.4,3.4,3])
para(r'Gate removal has little effect during the active thermal transient, while it loses the exact equilibrium-vanishing property. The safeguard also has a small nominal cost effect; its contribution is clearer under the deterministic 80 sin(0.7t) V patch corruption. The residual-damping ablation exposes omitted-mode excitation despite filtering. A finite, clipped trajectory in this ablation does not show that the high-order linear feedback is stable. The original three-mode synthesis counterexample remains positive at approximately 0.272 per second on a five-mode plant, demonstrating that a retained-coordinate inequality does not certify spillover stability.')

heading('6.3 Paired uncertainty and bounded commands',2)
rows=[]
for label in ['Oracle24','RawNN','FilteredNN']:
    for key,name in [('cost','Cost'),('tip_rms_m','Tip RMS'),('angle_rms_arcsec','Angle RMS')]:
        a=lhs['paired_summary'][label][key];ci=a['median_ratio_bootstrap95'];rows.append((label,name,f'{a["median_ratio"]:.4f}',f'[{ci[0]:.4f}, {ci[1]:.4f}]',f'{a["min_max"][1]:.4f}',f'{a["wins"]}/{lhs["N"]}'))
table('Table 10. Paired ratios against thermal LQR',['Controller','Metric','Median ratio','Bootstrap 95% interval','Worst ratio','Improved cases'],rows,[2.8,2.5,2.5,4.0,2.5,2.3])
figure('fig06_robustness','Figure 6. Distributions of paired filtered-NN-to-LQR ratios for all Latin-hypercube cases. The vertical dashed line is equal performance. Text gives the paired bootstrap interval for the median ratio. Unfavorable cases remain in the distributions.')
para(r'All paired differences, empirical percentile ranges, channel-level clipping fractions, longest saturation runs, RMS and peak inputs, and filter correction magnitudes are preserved in lhs_report.json. Table 11 summarizes extrema and medians across the uncertain design. Clipping records the candidate exceeding a bound before filtering; saturation records the applied command reaching a bound. A filter can reduce a clipped candidate, so these counts need not coincide.')
rows=[]
for label in ['LQR','Oracle24','RawNN','FilteredNN']:
    cases=[c[label] for c in lhs['results']];channels=np.array([c['channels'] for c in cases]);peak=channels[:,0];rms=channels[:,1];sat=channels[:,2];clip=channels[:,3];duration=channels[:,4]
    rows.append((label,f'{peak[:,0].max():.4f}',f'{peak[:,1:].max():.3f}',f'{sat.max()*100:.3f}',f'{clip.max()*100:.3f}',f'{duration.max():.4f}',f'{np.median([c["correction_rms_normalized"] for c in cases]):.3g}'))
table('Table 11. Uncertainty command statistics over every native sample',['Controller','Max torque N m','Max voltage V','Max channel saturation %','Max channel clipping %','Longest saturation s','Median correction RMS'],rows,[2.9,2.3,2.3,2.5,2.5,2.0,2.1])
para(r'These maxima combine different cases and channels and do not describe one simultaneous worst spacecraft. Nominal channel peaks, RMS values and fractions are supplied in Appendix B. Structured thickness and corner cases retain the same controller gains. They are reported separately because their parameter dependence differs from independent uniform margins; their numerical outcomes do not imply a physically identified tail probability.')
para(fr'Applied saturation occurs in {sum(any(np.array(c[label]["channels"])[2]>0) for c in lhs["results"] for label in ["LQR","Oracle24","RawNN","FilteredNN"])} of the {4*lhs["N"]} principal uncertain trajectories. Filter infeasibility totals {sum(c["FilteredNN"]["infeasible_count"] for c in lhs["results"]):.0f} samples in the filtered uncertainty cases. These statements concern the tested design; the sensor, omitted-damping and delayed-residual stress cases have different command behavior.')
rows=[]
for name,a in extra['structured'].items():rows.append((name,f'{a["cost"]:.5g}',f'{a["tip_rms_m"]*1e6:.2f}',f'{a["angle_rms_arcsec"]:.3f}',f'{a["infeasible_count"]:.0f}'))
table('Table 12. Correlated and corner scenarios',['Scenario and controller','Cost','Tip RMS µm','Angle RMS arcsec','Infeasible samples'],rows,[6.0,2.6,2.8,3.3,1.9])
para(r'Figure 7 distinguishes the intentionally retained thermal bend from dynamic tracking error and displays the nominal inequality evaluated at actual retained states. That residual is a check of the specified nominal formula, rather than a certificate for the full uncertain plant. Figure 8 plots both transition directions for the horizon settings tabulated above, including actuator effort.')
figure('fig07_equilibrium_certificate','Figure 7. Quasistatic thermal tip reference, total tip deflection and sampled nominal inequality residual at exit. The large static bend is retained by the objective. The residual uses nominal dynamics and actual retained states; it is not the derivative of the uncertain high order plant.')
figure('fig08_horizon','Figure 8. Horizon dependence of cost, tip RMS, attitude RMS and normalized actuator mean square. The zero-horizon endpoint is thermal LQR. Lines connect computed settings only.')

heading('6.4 Spatial validation and sampling limits',2)
meshes=fe['meshes']+refine['fine_FE'];rows=[]
for a in meshes:
    b=a['dynamic'];rows.append((a['actual_elements'],f'{b["LQR"]["cost"]:.6f}',f'{b["FilteredNN"]["cost"]:.6f}',f'{b["LQR"]["tip_rms_m"]*1e6:.3f}',f'{b["FilteredNN"]["tip_rms_m"]*1e6:.3f}'))
table('Table 13. Independent FE mesh refinement with 37 retained FE modes',['Elements','LQR cost','NN cost','LQR tip µm','NN tip µm'],rows,[2.4,3.5,3.5,3.6,3.6])
last=meshes[-1]['dynamic']['FilteredNN'];before=meshes[-2]['dynamic']['FilteredNN']
para(fr'On the two finest FE meshes, filtered-NN tip RMS changes from {before["tip_rms_m"]*1e6:.4f} to {last["tip_rms_m"]*1e6:.4f} µm, a relative change of {100*abs(last["tip_rms_m"]/before["tip_rms_m"]-1):.3f}%. This is mesh convergence at fixed modal retention and damping definition. It does not establish convergence of every distributed-parameter output. The independent FE result remains within the same broad performance scale as the analytic trial-basis model, despite different high-mode damping.')
rows=[]
for a in refine['modal']:
    b=a['metrics'];rows.append((a['n'],f'{1/eval_dt:.0f}',f'{a["sampled_linear_radius"]:.9f}',f'{b["cost"]:.5g}',f'{b["tip_rms_m"]*1e6:.3f}'))
for a in refine['modal65_time_refinement']:
    b=a['metrics'];rows.append((65,f'{1/a["dt"]:.0f}',f'{a["sampled_linear_radius"]:.9f}',f'{b["cost"]:.5g}',f'{b["tip_rms_m"]*1e6:.3f}'))
table('Table 14. Analytic basis and 65 mode time refinement',['Elastic modes','Updates per second','Linear spectral radius','Filtered NN cost','Tip RMS µm'],rows,[2.6,3.2,4,3.4,3.4])
para(r'The 65-mode sampled linear radius exceeding one at 6400 updates/s is an implementation failure, even if clipping keeps a finite simulation bounded. Its result is included and is excluded from convergence claims. The finer update intervals assess whether resolving and damping those modes restores a stable sampled comparison. Deployment must co-design bandwidth, mode retention and residual-rate delay; increasing plant order alone cannot validate a fixed digital controller.')
rows=[]
for nc,a in fe['spatial_thermal'].items():
    for label,b in a['dynamic'].items():rows.append((nc,label,f'{b["cost"]:.5f}',f'{b["tip_rms_m"]*1e6:.2f}',f'{b["angle_rms_arcsec"]:.3f}'))
table('Table 15. Nonuniform thermal field and reference integration',['Cells per face','Controller and reference','Cost','Tip RMS µm','Angle RMS arcsec'],rows,[2.6,5.0,3,3,3])
para(r'The scalar-reference controllers leave a large persistent tip error relative to the FE plant’s actual spatial equilibrium. Mean-temperature preview alone does not cure that shape error. The field-aware comparisons separate this reference mismatch from the benefit of forecast timing. Their added structural-reference and residual-feedforward calculations are model based and require more thermal information; they are not evidence that the original network generalizes to an arbitrary field. Mesh/cell refinement and the spatial heat-map in Figure 9 show the tested forcing and output sensitivity.')
figure('fig09_spatial_convergence','Figure 9. Independent FE mesh refinement, analytic modal refinement, and computed nonuniform two-face thermal difference. The unstable 65-mode result at 6400 updates/s is retained to expose the sampling limit. The thermal field is generated by prescribed flux gradient and shadow sweep.')

heading('6.5 Sensing delay and amplifier effects',2)
rows=[]
for category in ['sensing','delays','drivers']:
    for label,a in extra[category].items():rows.append((label,f'{a["cost"]:.5g}',f'{a["tip_rms_m"]*1e6:.2f}',f'{a["angle_rms_arcsec"]:.3f}',f'{a["true_margin_violation_fraction"]*100:.2f}'))
table('Table 16. Output feedback and implementation sensitivity',['Case','Cost','Tip RMS µm','Angle RMS arcsec','Nominal margin violations at true state %'],rows,[4.7,2.7,3,3.2,3])
para(r'Noise and bias levels are explicit engineering assumptions, not calibrated flight sensor specifications. The nominal filter can remain satisfied at an estimated state while its nominal expression at the true retained state is violated. A favorable output-feedback trajectory therefore does not extend the state-feedback Lyapunov proposition. The additional residual-port experiment tests noise in the otherwise ideal spillover damping measurement; its command-equivalent standard deviations are 5 mV and 50 mV.')
rows=[]
for a in observer_report['spectrum']:rows.append((f'{a["normalized_process_covariance"]:.0e}',f'{a["floquet_radius_5ms"]:.6f}',f'{a["equivalent_spectral_rate_s-1"]:.3f}'))
table('Table 16a. Observer process covariance and high order sampled spectrum',['Normalized process covariance per 1 ms','Five millisecond Floquet radius','Equivalent rate per second'],rows,[7,4.8,4.8])
para(fr'The initial observer setting is a failed deployment example: its 5 ms periodic sampled transition has radius greater than one, and the resulting costs are large in Table 16. Reducing normalized process covariance through $10^{{-8}}$ and $10^{{-10}}$ is still insufficient. The tested $10^{{-12}}$ setting yields radius {observer_report["spectrum"][-1]["floquet_radius_5ms"]:.6f}; it is the sole stable candidate among the four tested settings. This low-bandwidth estimator reduces high-mode measurement spillover, but its slower correction and initial-state sensitivity require further identification. The finite simulations in Table 16b test the same declared sensing conditions with this selected nominal design; no robust observer theorem is inferred.')
rows=[]
for label,a in observer_report['selected_cases'].items():rows.append((label,f'{a["cost"]:.5g}',f'{a["tip_rms_m"]*1e6:.2f}',f'{a["angle_rms_arcsec"]:.3f}',f'{a["true_margin_violation_fraction"]*100:.2f}'))
table('Table 16b. Selected slow observer with the same sensing assumptions',['Case','Cost','Tip RMS µm','Angle RMS arcsec','Nominal true state margin violations %'],rows,[4.7,2.7,3,3.2,3])
para(r'The two calibrated rows set every persistent bias to zero and temperature noise to 0.005 K, keeping 0.02 arcsec angle, 0.02 arcsec/s gyro and 5 nanostrain noise. This is an additional assumed calibration scenario, not a measured sensor specification. Comparing it with the biased rows quantifies thermal-reference accuracy separately from the stable observer choice. A precise equilibrium-coordinate controller can lose its micrometre tracking advantage through a small temperature-difference bias.')
rows=[]
for label,vv in extra['residual_sensor_noise'].items():
    a=vv['metrics'];rows.append((label,f'{vv["assumed_port_noise_std_Nm_per_Vs"]:.0e}',f'{a["cost"]:.5g}',f'{a["tip_rms_m"]*1e6:.2f}'))
table('Table 17. Separate residual strain rate sensor perturbations',['Controller and noise level','Port standard deviation N m per V s','Cost','Tip RMS µm'],rows,[5.2,4.3,3.5,3.6])
figure('fig10_sensing_delay','Figure 10. Paired output-feedback costs and core-command delay sensitivity. Temperature and strain biases persist through each case; low/high noise cases share their random draws across controllers. Delay labels use the native update interval.')

heading('6.6 Computation and actuator accounting',2)
para(fr'Warmed float64 CPU measurements use {bench["CPU"]}, {bench["platform"]}, Python {bench["Python"]} and NumPy {bench["NumPy"]}. All four feedforward methods use compiled Numba without fast-math and one BLAS thread. JIT compilation, mechanical exponential preparation, training and file access are excluded. Fifteen repetitions of 1000 common thermal states measure individual-call median, 99th percentile and amortized batch time. Forecast kernels are cached before timing, making the analytic comparator materially stronger than recomputing exponentials online. No embedded processor, interrupt jitter or hard real-time deadline is certified.')
rows=[]
for label,a in bench['timings'].items():rows.append((label,f'{a["median_per_call_us"]:.3f}',f'{a["p99_per_call_us"]:.3f}',f'{a["median_batch_amortized_us"]:.3f}'))
table('Table 18. Warm feedforward evaluation on the measured CPU',['Method','Median µs','99th percentile µs','Amortized µs'],rows,[5.5,3.7,3.7,3.7])
para(fr'The measured median analytic-to-neural per-call ratio is {bench["timings"]["CachedOracle24"]["median_per_call_us"]/bench["timings"]["NN"]["median_per_call_us"]:.1f}. Recomputing analytic preview at every {eval_dt*1e6:.2f} µs native interval would exceed that interval on this CPU. However, the oracle used for closed-loop comparison refreshes every 0.1 s and interpolates a slowly changing map, so analytic forecasting remains practical at that slower rate. The NN permits cheaper calls at a common rate; these measurements do not establish that learning is necessary for the desktop implementation or quantify an embedded processor benefit. The cubic polynomial is faster per call but less accurate in the tested closed loop.')
mem=bench['storage_and_operations']
para(fr'The neural parameters and output scales require {mem["NN_weights_and_scales_bytes"]} bytes and {mem["NN_dense_multiply_accumulates"]} dense multiply-accumulates plus 80 tanh calls per inference. The cached 24 s thermal kernel uses {mem["Oracle24_cached_thermal_kernel_bytes"]} bytes, with 1600 RK4 thermal stages and 200 thermal-acceleration evaluations per forecast. Polynomial coefficients require {mem["Polynomial_coefficients_bytes"]} bytes plus {mem["Polynomial_integer_exponent_bytes"]} bytes for the generic exponent implementation. These figures describe each feedforward map and exclude the common feedback, observer, safeguard, plant arrays and runtime overhead. The NN trades larger coefficient storage for fewer forecast operations; its necessity depends on the measured accuracy and closed-loop costs of simpler alternatives.')
para(fr'RBF centers, coefficients and inverse width occupy {mem["RBF_centers_coefficients_gamma_bytes"]} bytes and require 256 Gaussian exponential evaluations, feature distances and the output product. Its operation mix differs from 80 neural tanh calls, so parameter count alone is insufficient to rank execution time. Table 18 reports the measured implementation for each.')
figure('fig11_timing_ablation','Figure 11. Warm CPU timing and nominal component-ablation cost. The timing panel marks the 99th percentile above each median bar. All methods use compiled arithmetic and cached preparation. Removing residual damping is reported on a logarithmic cost axis.')
para(r'Electrical accounting assumes relative permittivity 1700, driver output resistance 100 ohms per ceramic and dielectric loss tangent 0.02. These are illustrative circuit/material values, not measured properties of an installed actuator. Each 0.30 m by 0.50 m, 0.30 mm ceramic has capacitance $C_p=\epsilon_0\epsilon_rA_p/h_p$; two ceramics per pair receive opposite voltage. Full native voltage traces determine current, resistive loss and wheel momentum change. A mean-removed finite-record spectral loss proxy assumes a frequency-independent loss tangent; it does not include nonlinear hysteresis, amplifier efficiency or DC leakage.')
eq(r'i_p=C_p\dot V_p,\quad E_R=2R_dC_p^2\int\sum_{j=1}^{3}\dot V_j^2dt,\quad E_C^{max}=\max_t C_p\sum_{j=1}^{3}V_j^2')
eq(r'\Delta H_w(t)=-\int_0^t\tau_h(s)\,ds')
eq(r'E_{diel}\approx2C_p\tan\delta\sum_{j=1}^{3}\int_{-\infty}^{\infty}|\omega|\,|\tilde V_j(\omega)|^2\frac{d\omega}{2\pi}')
rows=[]
for label,a in bench['electrical'].items():rows.append((label,f'{max(a["rms_current_each_patch_A"])*1000:.4f}',f'{a["six_ceramic_resistive_loss_J"]:.4g}',f'{a["six_ceramic_spectral_dielectric_proxy_J"]:.4g}',f'{a["maximum_stored_electrical_J"]:.4g}',f'{a["wheel_peak_abs_momentum_change_Nms"]:.4f}'))
table('Table 19. Illustrative six ceramic and wheel implications over 120 s',['Controller','Max RMS current mA','Resistive loss J','Dielectric proxy J','Peak stored energy J','Peak wheel momentum N m s'],rows,[2.8,2.9,2.6,2.7,2.8,2.8])
para(r'The sign of wheel momentum is opposite the hub torque; the table reports absolute excursion. A real wheel must also carry mission attitude momentum and desaturation margins. Bond-transfer efficiency, ceramic temperature dependence, fatigue, polarization, hysteresis and measured dielectric properties remain outside the present model. The tested first-order driver cases quantify one bandwidth limitation without constituting a complete amplifier design.')

heading('7 Limitations and spacecraft integration')
para(r'The contribution is a control integration evaluated on reproducible simulated benchmarks. Classical finite-preview theory supplies its nominal optimizer; no priority is claimed for that general solution. Constraints, filtering, residual damping, model mismatch and sampled execution change the implemented optimum. The conditional Lyapunov argument is not a barrier certificate, a robust nonlinear flutter guarantee or an infinite-dimensional stability proof.')
para(r'Independent FE and spatial thermal calculations improve model checks while retaining planar small-deflection bending. They omit three-axis maneuvers, torsion, hinge and SADA flexibility, attitude-dependent illumination, geometric shadowing, laminate asymmetry, material-temperature identification and nonlinear deformation. Field-aware reference calculations require temperature-field information whose observability and hardware measurement layout are not established here. The simulated observer assumes correct initial equilibrium and declared noise models; its parameters must be identified before hardware use.')
para(r'The uncertainty design is broader than the initial forty pairs but still depends on chosen marginal ranges and idealized correlations. Bootstrap intervals summarize the tested design. They do not imply flight reliability or calibrated manufacturing distributions. Neural sensitivity uses one initialization seed and a fixed synthetic envelope; extrapolation outside that envelope remains unqualified. CPU results concern the measured desktop implementation, with common controller costs excluded from feedforward timing. Electronics accounting is illustrative until material and drive parameters are measured.')

heading('8 Conclusions')
para(r'Moving thermal-equilibrium coordinates expose the disturbance rate that classical preview control can anticipate. A compact neural correction approximates the nonlinear thermal-to-preview map while retaining an analytic feedback backbone. The paired comparisons, horizon study and polynomial/RBF comparators quantify when forecast compression helps and when simpler choices suffice. The safeguard limits selected harmful learned commands but cannot repair unobserved high-mode dynamics or a wrong thermal reference. Independent FE, spatial heat and output-feedback tests show why sensing, reference reconstruction and digital bandwidth belong in the integration strategy. Exact optimality remains confined to the specified nominal unconstrained problem; the implemented spacecraft-control evidence is numerical and conditional.')

heading('Appendix A Reproducibility and archive')
para(r'The versioned package includes simulation and authoring scripts, the independent MATLAB check, all trained parameter sets, polynomial coefficients, fixed splits, Latin-hypercube design, per-case outcomes, native-step statistics, FE and spatial thermal histories, equation inventory and figures in PNG/PDF/SVG formats. Requirements, commands, seeds, environment details and SHA-256 hashes are recorded. Licensed publisher PDFs and third-party book illustrations are excluded from redistribution; source hashes and verified reference metadata retain provenance.')
para('The reproducibility package, version 2.0.0, is deposited in Zenodo at https://doi.org/10.5281/zenodo.23126822. It contains the source code, computed datasets, high-resolution and vector figures, editable manuscript, reference-verification records and SHA-256 manifest. Generated data, figures and documentation are licensed under CC BY 4.0; newly authored source code is licensed under MIT. The archive creator and contact is Caglar Uyulan (caglar.uyulan@ikcu.edu.tr). Publisher PDFs, third-party book illustrations and the supplied original manuscripts are excluded from redistribution.')
heading('Appendix B Nominal command statistics')
rows=[]
for name in ['LQR','Oracle24','RawNN','FilteredNN']:
    a=extra['nominal'][name+'_exit'];ch=np.array(a['channels'])
    for k,label in enumerate(['Hub torque','Patch 1','Patch 2','Patch 3']):rows.append((name,label,f'{ch[0,k]:.5g}',f'{ch[1,k]:.5g}',f'{100*ch[2,k]:.3f}',f'{100*ch[3,k]:.3f}',f'{ch[4,k]:.5f}'))
table('Table B1. Exit channel statistics from every native sample',['Controller','Channel','Peak','RMS','Saturated %','Clipped %','Longest saturation s'],rows,[2.9,2.7,2.3,2.3,2.3,2.3,1.8])
para(r'Torque peak/RMS units are N m and patch units are V. The complete entry table and all uncertain per-channel statistics are in the numerical report. Filter intervention RMS and maximum are normalized by channel bounds; they quantify the difference from the clipped candidate, separately from clipping itself.')

heading('References')
records=json.loads((DATA/'references/crossref_verified.json').read_text(encoding='utf-8'))
def format_record(a,i):
    authors=[]
    for person in a['authors']:
        given=person.get('given','').replace('\ufffd','');initials=''.join(word[0] for word in re.findall(r'[A-Za-zÀ-ž]+',given))
        if person['family'] in ['Becerikli','Samad']:initials={'Becerikli':'Y','Samad':'T'}[person['family']]
        authors.append(person['family']+' '+initials)
    date=a.get('published_print') or a['published'];year=date['date-parts'][0][0]
    container=(a['container'] or [a['publisher']])[0];details=str(year)
    if a['volume']:details+=';'+a['volume']
    if a['issue']:details+='('+a['issue']+')'
    if a['pages'] or a['article_number']:details+=':'+(a['pages'] or a['article_number'])
    return f'[{i}] '+', '.join(authors)+'. '+a['title'][0]+'. '+container+'. '+details+'. https://doi.org/'+a['resolved_doi']+'.'
refs=[format_record(a,i+1) for i,a in enumerate(records[:14])]
refs.extend([
'[15] Nakamura-Zimmerer T. QRnet software repository. https://github.com/Tenavi/QRnet. Accessed 3 October 2026. Software reference cited by URL.',
'[16] Yavuz MT, Uyulan Ç, Acarer S. Intelligent optimal control of thermally induced vibrations in satellite components under solar heat flux. Supplied UHUK 2026 abstract manuscript in Turkish. DOI not supplied with the source; publication status is not established here.',
'[17] National Institute of Standards and Technology. CODATA value of the Stefan Boltzmann constant. 2022 recommended values. https://physics.nist.gov/cgi-bin/cuu/Value?sigma. Accessed 3 October 2026. Web reference without a DOI.'
])
refs.extend(format_record(a,i+18) for i,a in enumerate(records[14:]))
for text in refs:
    p=doc.add_paragraph(text);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.;p.paragraph_format.left_indent=Cm(.55);p.paragraph_format.first_line_indent=Cm(-.55)
    for rr in p.runs:rr.font.size=Pt(10)
text_log.extend(refs)
filename='Thermal_Preview_Learning_and_Lyapunov_Safeguarding_Revised'
doc.save(OUT/(filename+'.docx'))
(DATA/'equation_inventory.json').write_text(json.dumps(equations,indent=2),encoding='utf-8')
(DATA/'manuscript_text.txt').write_text('\n\n'.join(text_log),encoding='utf-8')
with zipfile.ZipFile(OUT/(filename+'.docx')) as zz:
    xml=etree.fromstring(zz.read('word/document.xml'));ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    counts={'display_equations':len(equations),'OMML_math_objects':len(xml.xpath('.//m:oMath',namespaces=ns)),'word_count_estimate':len(re.findall(r'\b\w+\b',' '.join(text_log))),'references':len(refs),'figures':11,'tables':len(doc.tables)}
    assert counts['OMML_math_objects']>=len(equations)
(DATA/'document_structure_audit.json').write_text(json.dumps(counts,indent=2),encoding='utf-8')
print(json.dumps(counts,indent=2),flush=True)

